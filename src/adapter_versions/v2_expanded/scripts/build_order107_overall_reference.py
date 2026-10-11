#!/usr/bin/env python3
"""Build the Order 107 evaluator reference from the source ECS PDFs."""

from __future__ import annotations

import argparse
from pathlib import Path

import fitz
from openpyxl import load_workbook


YEARS = range(2015, 2024)
COLUMNS = ("CityU", "HKBU", "LU", "CUHK", "HKIEd (EduHK)", "PolyU", "HKUST", "HKU")


def only_cached(cache_dir: Path, suffix: str) -> Path:
    matches = list(cache_dir.glob(f"*_{suffix}"))
    if len(matches) != 1:
        raise ValueError(f"expected one cached {suffix}, found {len(matches)}")
    return matches[0]


def overall_rates(pdf_path: Path, columns: tuple[str, ...]) -> list[float]:
    with fitz.open(pdf_path) as document:
        if document.page_count != 1:
            raise ValueError(f"expected one-page ECS report: {pdf_path}")
        tables = document[0].find_tables().tables
        if len(tables) != 1:
            raise ValueError(f"expected one ECS table: {pdf_path}")
        rows = tables[0].extract()
    totals = [row for row in rows if row and row[0] == "Total"]
    if len(totals) != 1:
        raise ValueError(f"expected one overall Total row: {pdf_path}")
    headers, measures, units = rows[:3]
    total = totals[0]
    rates = []
    for column in columns:
        aliases = {column}
        if column == "HKIEd (EduHK)":
            aliases = {"HKIEd", "EdUHK"}
        starts = [index for index, label in enumerate(headers) if label in aliases]
        if len(starts) != 1:
            raise ValueError(f"missing university header {column}: {pdf_path}")
        start = starts[0]
        if (measures[start], units[start], measures[start + 2], units[start + 2]) != (
            "Application", "No", "Supported", "No"
        ):
            raise ValueError(f"unexpected count columns for {column}: {pdf_path}")
        applications = int(total[start].replace(",", ""))
        supported = int(total[start + 2].replace(",", ""))
        if not 0 <= supported <= applications:
            raise ValueError(f"invalid counts for {column}: {pdf_path}")
        rates.append(supported / applications if applications else 0.0)
    return rates


def build(cache_dir: Path, output: Path) -> None:
    workbook = load_workbook(only_cached(cache_dir, "supported_rate.xlsx"))
    sheet = workbook["Sheet1"]
    if tuple(cell.value for cell in sheet[1]) != ("Year", *COLUMNS):
        raise ValueError("unexpected target workbook headers")
    for row_index, year in enumerate(YEARS, start=2):
        if sheet.cell(row_index, 1).value != year:
            raise ValueError(f"unexpected year in row {row_index}")
        if any(sheet.cell(row_index, col).value is not None for col in range(2, 10)):
            raise ValueError(f"target row {row_index} is not blank")
        pdf = only_cached(cache_dir, f"ecs{year % 100:02d}.pdf")
        for col_index, rate in enumerate(overall_rates(pdf, COLUMNS), start=2):
            sheet.cell(row_index, col_index, rate)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cache-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.cache_dir, args.output)


if __name__ == "__main__":
    main()
